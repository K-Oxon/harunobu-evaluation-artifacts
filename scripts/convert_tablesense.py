"""Convert TableSense XLS files with LibreOffice and validate them with openpyxl. Writes XLSX files under data/tablesense/xlsx/ and a report under results/layout/."""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from scripts import paths
from scripts.ir import IRBBox
from scripts.loaders.tablesense import file_index, load_tablesense, raw_relpath

_SOFFICE = "soffice"
_BATCH_SIZE = 40
_BATCH_TIMEOUT_BASE = 60
_BATCH_TIMEOUT_PER_FILE = 10
_SINGLE_TIMEOUT = 120


def _nonempty(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def _bbox_ref(bbox: IRBBox) -> str:
    return (
        f"{get_column_letter(bbox.first_col + 1)}{bbox.first_row + 1}:"
        f"{get_column_letter(bbox.last_col + 1)}{bbox.last_row + 1}"
    )


@dataclass
class FileResult:

    status: str = "ok"
    corpus: str = ""
    relpath: str = ""
    sheets_nonempty: dict[str, int] = field(default_factory=dict)
    sheets_missing: list[str] = field(default_factory=list)
    empty_ranges: list[str] = field(default_factory=list)
    empty_ranges_in_source: list[str] = field(default_factory=list)
    xls_sheets_missing: list[str] = field(default_factory=list)
    note: str = ""

    def as_dict(self) -> dict:
        d: dict = {"status": self.status, "corpus": self.corpus, "relpath": self.relpath}
        if self.status == "ok":
            d["sheets_nonempty"] = dict(sorted(self.sheets_nonempty.items()))
        for key in ("sheets_missing", "empty_ranges", "empty_ranges_in_source", "xls_sheets_missing"):
            v = getattr(self, key)
            if v:
                d[key] = sorted(v)
        if self.note:
            d["note"] = self.note
        return d


def _xls_path(relpath: str) -> Path:
    path_field, file_name = relpath.rsplit("/", 1)
    return paths.TABLESENSE_RAW_DIR / raw_relpath(path_field, file_name)


def _soffice_convert(xls_paths: list[Path], outdir: Path, timeout: int) -> None:
    with tempfile.TemporaryDirectory(prefix="lo_profile_") as profile:
        subprocess.run(
            [
                _SOFFICE,
                "--headless",
                "--norestore",
                f"-env:UserInstallation=file://{profile}",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(outdir),
                *[str(p) for p in xls_paths],
            ],
            check=True,
            capture_output=True,
            timeout=timeout,
        )


def _convert_batch(batch: list[tuple[str, Path, Path]], results: dict[str, FileResult]) -> None:
    with tempfile.TemporaryDirectory(prefix="ts_convert_") as tmp:
        tmpdir = Path(tmp)
        timeout = _BATCH_TIMEOUT_BASE + _BATCH_TIMEOUT_PER_FILE * len(batch)
        batch_err: str | None = None
        try:
            _soffice_convert([x for _, x, _ in batch], tmpdir, timeout)
        except subprocess.TimeoutExpired:
            batch_err = "timeout"
        except subprocess.CalledProcessError as exc:
            batch_err = f"exit={exc.returncode}"

        for file_name, xls_path, final_path in batch:
            produced = tmpdir / (xls_path.stem + ".xlsx")
            if produced.exists():
                final_path.parent.mkdir(parents=True, exist_ok=True)
                shutil.move(str(produced), str(final_path))
                continue

            try:
                with tempfile.TemporaryDirectory(prefix="ts_retry_") as tmp2:
                    _soffice_convert([xls_path], Path(tmp2), _SINGLE_TIMEOUT)
                    produced2 = Path(tmp2) / (xls_path.stem + ".xlsx")
                    if produced2.exists():
                        final_path.parent.mkdir(parents=True, exist_ok=True)
                        shutil.move(str(produced2), str(final_path))
                        continue
                    results[file_name].status = "convert_failed"
                    results[file_name].note = f"soffice 出力なし(バッチ: {batch_err or 'ok'})"
            except subprocess.TimeoutExpired:
                results[file_name].status = "timeout"
            except subprocess.CalledProcessError as exc:
                results[file_name].status = "convert_failed"
                results[file_name].note = f"soffice exit={exc.returncode}"


def _validate_file(
    result: FileResult,
    xlsx_path: Path,
    xls_path: Path,
    gt_sheets: dict[str, list[IRBBox]],
) -> None:
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    except Exception as exc:
        result.status = "convert_failed"
        result.note = f"openpyxl 読込失敗: {type(exc).__name__}"
        return

    try:
        wb_sheets = set(wb.sheetnames)
        for sheet_name, bboxes in gt_sheets.items():
            if sheet_name not in wb_sheets:
                result.sheets_missing.append(sheet_name)
                continue
            ws = wb[sheet_name]
            nonempty_total = 0
            occupied: set[tuple[int, int]] = set()
            for row in ws.iter_rows():
                for cell in row:
                    if _nonempty(cell.value):
                        nonempty_total += 1
                        occupied.add((cell.row - 1, cell.column - 1))
            result.sheets_nonempty[sheet_name] = nonempty_total
            for bbox in bboxes:
                hit = any(
                    (r, c) in occupied
                    for r in range(bbox.first_row, bbox.last_row + 1)
                    for c in range(bbox.first_col, bbox.last_col + 1)
                )
                if not hit:
                    label = f"{sheet_name}!{_bbox_ref(bbox)}"
                    if _range_empty_in_source(xls_path, sheet_name, bbox):
                        result.empty_ranges_in_source.append(label)
                    else:
                        result.empty_ranges.append(label)
    finally:
        wb.close()

    if result.sheets_missing:
        result.status = "sheet_mismatch"
    elif result.empty_ranges:
        result.status = "empty_range"


def _range_empty_in_source(xls_path: Path, sheet_name: str, bbox: IRBBox) -> bool:
    import xlrd

    try:
        book = xlrd.open_workbook(str(xls_path), on_demand=True)
        try:
            sheet = book.sheet_by_name(sheet_name)
        except xlrd.XLRDError:
            return False
        for r in range(bbox.first_row, min(bbox.last_row + 1, sheet.nrows)):
            for c in range(bbox.first_col, min(bbox.last_col + 1, sheet.ncols)):
                if _nonempty(sheet.cell_value(r, c)):
                    return False
        return True
    except Exception:
        return False


def _xlrd_sheet_check(xls_path: Path, annotated_sheets: list[str]) -> list[str]:
    import xlrd

    try:
        book = xlrd.open_workbook(str(xls_path), on_demand=True)
        names = set(book.sheet_names())
        return [s for s in annotated_sheets if s not in names]
    except Exception:
        return []


def run(limit: int | None = None, force: bool = False) -> dict:
    index = file_index()
    docs = load_tablesense()
    names = sorted(index.keys())
    if limit:
        names = names[:limit]

    results: dict[str, FileResult] = {}
    to_convert: list[tuple[str, Path, Path]] = []


    for name in names:
        info = index[name]
        result = FileResult(corpus=info.corpus, relpath=info.relpath)
        results[name] = result
        xls_path = _xls_path(info.relpath)
        if not xls_path.exists():
            result.status = "missing_xls"
            continue
        final_path = paths.TABLESENSE_XLSX_DIR / info.relpath
        if force or not final_path.exists():
            to_convert.append((name, xls_path, final_path))
    n_missing = sum(1 for r in results.values() if r.status == "missing_xls")
    print(f"resolve: {len(names)} 件中 missing_xls={n_missing}, 要変換={len(to_convert)}")


    batches: list[list[tuple[str, Path, Path]]] = []
    current: list[tuple[str, Path, Path]] = []
    stems: set[str] = set()
    for item in to_convert:
        stem = item[1].stem
        if len(current) >= _BATCH_SIZE or stem in stems:
            batches.append(current)
            current, stems = [], set()
        current.append(item)
        stems.add(stem)
    if current:
        batches.append(current)
    for i, batch in enumerate(batches):
        _convert_batch(batch, results)
        done = sum(len(b) for b in batches[: i + 1])
        print(f"convert: batch {i + 1}/{len(batches)} 済 ({done}/{len(to_convert)})", flush=True)


    for name in names:
        result = results[name]
        if result.status != "ok":
            continue
        xlsx_path = paths.TABLESENSE_XLSX_DIR / result.relpath
        if not xlsx_path.exists():
            result.status = "convert_failed"
            result.note = "変換済み xlsx が存在しない"
            continue
        xls_path = _xls_path(result.relpath)
        gt_sheets = {s.sheet_name: [t.bbox for t in s.tables] for s in docs[name].sheets}
        _validate_file(result, xlsx_path, xls_path, gt_sheets)
        result.xls_sheets_missing = _xlrd_sheet_check(xls_path, list(gt_sheets.keys()))


    by_status: dict[str, int] = {}
    for r in results.values():
        by_status[r.status] = by_status.get(r.status, 0) + 1
    payload = {
        "summary": {
            "total": len(names),
            "by_status": dict(sorted(by_status.items())),
            "n_empty_ranges_in_source": sum(len(r.empty_ranges_in_source) for r in results.values()),
            "soffice_version": _soffice_version(),
        },
        "files": {name: results[name].as_dict() for name in sorted(results)},
    }
    return payload


def _soffice_version() -> str:
    try:
        out = subprocess.run([_SOFFICE, "--version"], capture_output=True, text=True, timeout=30)
        return out.stdout.strip().splitlines()[0] if out.stdout else "unknown"
    except Exception:
        return "unknown"


def load_conversion_status(path: Path | None = None) -> dict[str, str]:
    path = path or paths.TABLESENSE_CONVERSION_JSON
    payload = json.loads(path.read_text(encoding="utf-8"))
    return {name: rec["status"] for name, rec in payload["files"].items()}


def main() -> int:
    parser = argparse.ArgumentParser(description="TableSense .xls → .xlsx 変換(resolve/convert/validate)")
    parser.add_argument("--limit", type=int, default=None, help="先頭 N 件のみ(スモーク用。JSON は書かない)")
    parser.add_argument("--force", action="store_true", help="変換済みも再変換する")
    args = parser.parse_args()

    payload = run(limit=args.limit, force=args.force)
    print(json.dumps(payload["summary"], ensure_ascii=False, indent=2))

    if args.limit is None:
        paths.TABLESENSE_CONVERSION_JSON.parent.mkdir(parents=True, exist_ok=True)
        paths.TABLESENSE_CONVERSION_JSON.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        print(f"→ {paths.TABLESENSE_CONVERSION_JSON}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
