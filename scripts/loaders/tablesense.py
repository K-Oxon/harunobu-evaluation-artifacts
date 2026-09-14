"""Load TableSense table ranges into the evaluation data model. Depends on the TableSense annotation TSV and returns in-memory documents."""

from __future__ import annotations

import csv
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

from openpyxl.utils import range_boundaries

from scripts import paths
from scripts.ir import IRBBox, IRDocument, IRSheet, IRTable


@dataclass(frozen=True)
class TSFileInfo:

    file_name: str
    relpath: str
    corpus: str
    split: str
    n_sheets: int
    n_tables: int


def a1_range_to_bbox(ref: str) -> IRBBox:
    min_col, min_row, max_col, max_row = range_boundaries(ref.strip())
    return IRBBox(
        first_row=min_row - 1,
        first_col=min_col - 1,
        last_row=max_row - 1,
        last_col=max_col - 1,
    )


def xlsx_relpath(path_field: str, file_name: str) -> str:
    return path_field.replace("\\", "/") + "/" + file_name


def raw_relpath(path_field: str, file_name: str) -> str:
    parts = path_field.replace("\\", "/").split("/")
    if len(parts) >= 2 and parts[0] == parts[1]:
        parts = parts[1:]
    stem = Path(file_name).stem
    return "/".join(parts) + "/" + stem + ".xls"


class _SheetAcc:

    def __init__(self, sheet_name: str, sheet_index: int) -> None:
        self.sheet_name = sheet_name
        self.sheet_index = sheet_index
        self.tables: list[IRTable] = []
        self._seen: set[tuple[int, int, int, int]] = set()

    def add_bbox(self, bbox: IRBBox) -> None:
        key = (bbox.first_row, bbox.first_col, bbox.last_row, bbox.last_col)
        if key in self._seen:
            return
        self._seen.add(key)
        self.tables.append(IRTable(bbox=bbox))


class _FileAcc:

    def __init__(self, path_field: str, split: str) -> None:
        self.path_field = path_field
        self.split = split
        self.by_sheet: dict[str, _SheetAcc] = {}

    def sheet(self, sheet_name: str) -> _SheetAcc:
        acc = self.by_sheet.get(sheet_name)
        if acc is None:
            acc = _SheetAcc(sheet_name, sheet_index=len(self.by_sheet))
            self.by_sheet[sheet_name] = acc
        return acc


def _parse(annotations: Path) -> dict[str, _FileAcc]:
    store: dict[str, _FileAcc] = {}
    with open(annotations, encoding="utf-8", newline="") as f:
        for row in csv.reader(f, delimiter="\t"):
            if len(row) < 5 or not row[0].strip():
                continue



            file_name = row[0].strip()
            sheet_name = row[1]
            split = row[2].strip()
            path_field = row[3].strip()
            ranges = [r for r in (c.strip() for c in row[4:]) if r]

            acc = store.get(file_name)
            if acc is None:
                acc = _FileAcc(path_field, split)
                store[file_name] = acc
            elif acc.path_field != path_field:
                raise ValueError(
                    f"FileName↔path の 1:1 前提が崩れています: {file_name!r} が "
                    f"{acc.path_field!r} と {path_field!r} の両方に出現"
                )

            sheet = acc.sheet(sheet_name)
            for ref in ranges:
                sheet.add_bbox(a1_range_to_bbox(ref))
    return store


def load_tablesense(
    file_names: Iterable[str] | None = None,
    *,
    annotations: Path | None = None,
) -> dict[str, IRDocument]:
    annotations = annotations or paths.TABLESENSE_TSV
    targets = set(file_names) if file_names is not None else None

    result: dict[str, IRDocument] = {}
    for file_name, acc in _parse(annotations).items():
        if targets is not None and file_name not in targets:
            continue
        sheets = [
            IRSheet(
                sheet_name=s.sheet_name,
                sheet_index=s.sheet_index,
                tables=s.tables,
                cell_roles=[],
            )
            for s in acc.by_sheet.values()
        ]
        result[file_name] = IRDocument(file_name=file_name, source="tablesense", sheets=sheets)
    return result


def file_index(annotations: Path | None = None) -> dict[str, TSFileInfo]:
    annotations = annotations or paths.TABLESENSE_TSV
    index: dict[str, TSFileInfo] = {}
    for file_name, acc in _parse(annotations).items():
        corpus = acc.path_field.replace("\\", "/").split("/")[0]
        index[file_name] = TSFileInfo(
            file_name=file_name,
            relpath=xlsx_relpath(acc.path_field, file_name),
            corpus=corpus,
            split=acc.split,
            n_sheets=len(acc.by_sheet),
            n_tables=sum(len(s.tables) for s in acc.by_sheet.values()),
        )
    return index


def annotated_file_names(annotations: Path | None = None) -> list[str]:
    return sorted(file_index(annotations).keys())
